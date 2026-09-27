from io import BytesIO

import frappe
from docx import Document
from docx.oxml.ns import qn
from frappe.tests.utils import FrappeTestCase

from docstudio.engine.registry import build_model
from docstudio.export import richtext
from docstudio.export.docx_writer import build_docx
from docstudio.tests.utils import TAMIL, TEST_DT, make_record


def open_docx(content):
	return Document(BytesIO(content))


def all_text(wdoc):
	parts = [p.text for p in wdoc.paragraphs]

	def walk(tables):
		for t in tables:
			for row in t.rows:
				for cell in row.cells:
					parts.extend(p.text for p in cell.paragraphs)
					walk(cell.tables)

	walk(wdoc.tables)
	return "\n".join(parts)


class TestDocxExport(FrappeTestCase):
	@classmethod
	def setUpClass(cls):
		super().setUpClass()
		cls.doc = make_record(customer_name=TAMIL)

	def export(self, patch=None, theme=None):
		model = build_model(
			frappe.get_doc(TEST_DT, self.doc.name), theme_name=theme, patch=patch, preview=False
		)
		return open_docx(build_docx(model))

	def test_content(self):
		text = all_text(self.export())
		for expected in (
			TEST_DT.upper(),
			"Acme Test Traders",
			self.doc.name,
			"Cement",
			"Sand",
			"4,371.00",
			"For Acme Test Traders",
			"Authorized Signatory",
			"Deliver before noon.",
		):
			self.assertIn(expected, text)
		self.assertNotIn("SHOULD-NOT-SHOW", text)
		self.assertNotIn("ALSO-HIDDEN", text)

	def test_items_table(self):
		wdoc = self.export()
		items = next(t for t in wdoc.tables if t.rows[0].cells[0].text == "S.No")
		self.assertEqual([c.text for c in items.rows[0].cells], ["S.No", "Item", "Qty", "Rate", "Amount"])
		self.assertEqual(len(items.rows), 3)
		self.assertEqual(items.rows[1].cells[0].text, "1")
		shd = items.rows[0].cells[0]._tc.tcPr.find(qn("w:shd"))
		self.assertIsNotNone(shd)
		self.assertEqual(shd.get(qn("w:fill")), "D9E2F3")  # Business theme header shading
		# Header row repeats on every page.
		self.assertIsNotNone(items.rows[0]._tr.trPr.find(qn("w:tblHeader")))
		# Amount column right aligned.
		self.assertEqual(items.rows[1].cells[4].paragraphs[0].alignment, 2)

	def test_tamil_text_and_complex_script_fonts(self):
		wdoc = self.export()
		self.assertIn(TAMIL, all_text(wdoc))
		run = next(r for p in wdoc.paragraphs for r in p.runs if r.text)
		rfonts = run._element.rPr.find(qn("w:rFonts"))
		self.assertEqual(rfonts.get(qn("w:ascii")), "Arial")
		self.assertEqual(rfonts.get(qn("w:cs")), "Nirmala UI")
		normal = wdoc.styles["Normal"].element.rPr.find(qn("w:rFonts"))
		self.assertEqual(normal.get(qn("w:cs")), "Nirmala UI")

	def test_noto_theme_uses_noto_for_complex_scripts(self):
		if not frappe.db.exists("DocStudio Theme", "Test Tamil"):
			frappe.get_doc(
				{"doctype": "DocStudio Theme", "theme_name": "Test Tamil", "font_family": "Noto Sans Tamil"}
			).insert()
		wdoc = self.export(theme="Test Tamil")
		normal = wdoc.styles["Normal"].element.rPr.find(qn("w:rFonts"))
		self.assertEqual(normal.get(qn("w:cs")), "Noto Sans Tamil")
		self.assertEqual(normal.get(qn("w:ascii")), "Noto Sans Tamil")

	def test_page_setup(self):
		sec = self.export().sections[0]
		self.assertAlmostEqual(sec.page_width.mm, 210, places=0)
		self.assertAlmostEqual(sec.page_height.mm, 297, places=0)
		self.assertAlmostEqual(sec.left_margin.mm, 18, places=0)

	def test_edits_and_styles_reach_the_file(self):
		patch = {
			"fields": {"reference_no": "REF-EDITED"},
			"overrides": {"title:heading": "TAX INVOICE"},
			"styles": {"f:reference_no": {"bold": True, "color": "#AA0000"}},
			"layout": {"hidden": ["s:signature"]},
		}
		wdoc = self.export(patch)
		text = all_text(wdoc)
		self.assertIn("REF-EDITED", text)
		self.assertIn("TAX INVOICE", text)
		self.assertNotIn("Authorized Signatory", text)
		runs = [
			r
			for t in wdoc.tables
			for row in t.rows
			for c in row.cells
			for inner in c.tables
			for irow in inner.rows
			for ic in irow.cells
			for p in ic.paragraphs
			for r in p.runs
			if r.text == "REF-EDITED"
		] + [
			r
			for t in wdoc.tables
			for row in t.rows
			for c in row.cells
			for p in c.paragraphs
			for r in p.runs
			if r.text == "REF-EDITED"
		]
		self.assertTrue(runs)
		self.assertTrue(runs[0].bold)
		self.assertEqual(str(runs[0].font.color.rgb), "AA0000")

	def test_no_editor_artifacts(self):
		xml = self.export()._element.xml
		self.assertNotIn("ds-", xml)
		self.assertNotIn('w:val="dotted"', xml)

	def test_richtext_parser(self):
		blocks = richtext.parse(
			'<p class="ql-align-center">A <strong>b</strong> <em>c</em></p><ol><li>one</li><li>two</li></ol>'
		)
		self.assertEqual(blocks[0]["align"], "center")
		self.assertEqual(blocks[0]["runs"][1], ("b", {"bold": True}))
		self.assertEqual(blocks[2]["list"], ("ol", 2, 1))

	def test_richtext_css_styles(self):
		blocks = richtext.parse(
			'<p><span style="font-weight: bold; color: rgb(170, 0, 0)">x</span><span style="background-color:#ffffff">y</span></p>'
		)
		self.assertEqual(blocks[0]["runs"][0], ("x", {"bold": True, "color": "#AA0000"}))
		self.assertEqual(blocks[0]["runs"][1], ("y", {}))

	def test_clean_html_strips_external_css(self):
		from docstudio.engine.formatting import clean_html

		out = clean_html(
			'<span style="color: red; background-image: url(http://x/y.png)">a</span><script>1</script>'
		)
		self.assertNotIn("url(", out)
		self.assertNotIn("<script", out)
		self.assertIn("color: red", out)
