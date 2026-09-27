"""Write a DocModel to .docx with python-docx."""

import os
from io import BytesIO

import frappe
from frappe.utils import cstr, flt

from docstudio.engine.theme import PAGE_SIZES_MM
from docstudio.export import richtext

LABEL_COLOR = "#555555"
TAMIL_FALLBACK_FONT = "Nirmala UI"


def build_docx(model):
	"""Return .docx bytes for a model."""
	return DocxWriter(model).build()


def _rgb(hex_color):
	from docx.shared import RGBColor

	hex_color = cstr(hex_color).lstrip("#")
	if len(hex_color) != 6:
		return None
	return RGBColor.from_string(hex_color.upper())


def _hex_attr(hex_color):
	return cstr(hex_color).lstrip("#").upper() or "auto"


class DocxWriter:
	def __init__(self, model):
		self.model = model
		self.theme = model["theme"]
		self.doc = None

	def build(self):
		self.setup()
		self.write_sections()
		return self.save()

	def save(self):
		buf = BytesIO()
		self.doc.save(buf)
		return buf.getvalue()

	def setup(self):
		from docx import Document
		from docx.shared import Mm

		self.doc = Document()
		t = self.theme
		width, height = PAGE_SIZES_MM.get(t["page_size"], PAGE_SIZES_MM["A4"])
		for sec in self.doc.sections:
			sec.page_width, sec.page_height = Mm(width), Mm(height)
			sec.top_margin, sec.bottom_margin = Mm(t["margin_top"]), Mm(t["margin_bottom"])
			sec.left_margin, sec.right_margin = Mm(t["margin_left"]), Mm(t["margin_right"])
		self.content_mm = width - flt(t["margin_left"]) - flt(t["margin_right"])
		self._setup_styles()

	@property
	def cs_font(self):
		family = self.theme["font_family"]
		return family if family.startswith("Noto") else TAMIL_FALLBACK_FONT

	def _setup_styles(self):
		from docx.shared import Pt

		normal = self.doc.styles["Normal"]
		normal.font.name = self.theme["font_family"]
		normal.font.size = Pt(self.theme["base_font_size"])
		color = _rgb(self.theme["text_color"])
		if color:
			normal.font.color.rgb = color
		self._set_fonts(normal.element.get_or_add_rPr(), self.theme["base_font_size"])
		pf = normal.paragraph_format
		pf.space_before = Pt(0)
		pf.space_after = Pt(2)
		pf.line_spacing = flt(self.theme["line_spacing"]) or 1.15

	def _set_fonts(self, rPr, size_pt=None, bold=False, italic=False):
		"""Set every rFonts slot, including w:cs (complex scripts such as Tamil)."""
		from docx.oxml.ns import qn

		rFonts = rPr.get_or_add_rFonts()
		family = self.theme["font_family"]
		for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
			rFonts.set(qn(attr), family)
		rFonts.set(qn("w:cs"), self.cs_font)
		if bold:
			_ensure(rPr, "w:bCs")
		if italic:
			_ensure(rPr, "w:iCs")
		if size_pt:
			_ensure(rPr, "w:szCs").set(qn("w:val"), str(round(flt(size_pt) * 2)))
		_ensure(rPr, "w:lang").set(qn("w:bidi"), "ta-IN")

	def run(self, paragraph, text, style=None):
		from docx.shared import Pt

		style = style or {}
		parts = cstr(text).split("\n")
		first = None
		for i, part in enumerate(parts):
			r = paragraph.add_run(part)
			if i < len(parts) - 1:
				r.add_break()
			first = first or r
			if style.get("bold"):
				r.bold = True
			if style.get("italic"):
				r.italic = True
			if style.get("underline"):
				r.underline = True
			size = style.get("size")
			if size:
				r.font.size = Pt(size)
			color = _rgb(style.get("color")) if style.get("color") else None
			if color:
				r.font.color.rgb = color
			self._set_fonts(
				r._element.get_or_add_rPr(),
				size or self.theme["base_font_size"],
				bold=style.get("bold"),
				italic=style.get("italic"),
			)
		return first

	def align(self, paragraph, align):
		from docx.enum.text import WD_ALIGN_PARAGRAPH

		mapping = {
			"left": WD_ALIGN_PARAGRAPH.LEFT,
			"center": WD_ALIGN_PARAGRAPH.CENTER,
			"right": WD_ALIGN_PARAGRAPH.RIGHT,
			"justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
		}
		if align in mapping:
			paragraph.alignment = mapping[align]

	def paragraph(self, container=None, reuse=False):
		container = container or self.doc
		if (
			reuse
			and container.paragraphs
			and not container.paragraphs[-1].text
			and not container.paragraphs[-1].runs
		):
			return container.paragraphs[-1]
		return container.add_paragraph()

	def page_break(self):
		from docx.enum.text import WD_BREAK

		self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

	def node_style(self, node, extra=None):
		style = dict(extra or {})
		style.update(node.get("style") or {})
		return style

	def write_node(self, container, node, extra_style=None, reuse=True, align=None):
		"""Write a node's display value into a document/cell; handles rich text and newlines."""
		if node.get("hidden"):
			return
		style = self.node_style(node, extra_style)
		if node.get("html"):
			self.write_rich(container, node["display"], style, reuse=reuse)
			return
		p = self.paragraph(container, reuse=reuse)
		self.align(p, style.get("align") or align or node.get("align"))
		self.run(p, node["display"], style)

	def write_rich(self, container, html, base_style=None, reuse=True):
		base_style = base_style or {}
		blocks = richtext.parse(html)
		for i, block in enumerate(blocks):
			p = self.paragraph(container, reuse=reuse and i == 0)
			self.align(p, block["align"] or base_style.get("align"))
			if block["list"]:
				kind, number, depth = block["list"]
				p.paragraph_format.left_indent = self._mm(5 * depth)
				p.paragraph_format.first_line_indent = self._mm(-4)
				self.run(p, f"{number}.\t" if kind == "ol" else "•\t", base_style)
			heading_size = None
			if block["heading"]:
				heading_size = max(
					self.theme["base_font_size"] + 6 - block["heading"], self.theme["base_font_size"]
				)
			for text, run_style in block["runs"]:
				style = {**base_style, **run_style}
				if heading_size and "size" not in base_style:
					style["size"] = heading_size
				self.run(p, text, style)

	@staticmethod
	def _mm(value):
		from docx.shared import Mm

		return Mm(value)

	def table(self, container, rows, cols, widths_mm):
		from docx.oxml.ns import qn

		table = container.add_table(rows=rows, cols=cols)
		table.autofit = False
		_ensure(table._tbl.tblPr, "w:tblLayout").set(qn("w:type"), "fixed")
		for i, w in enumerate(widths_mm):
			table.columns[i].width = self._mm(w)
			for cell in table.columns[i].cells:
				cell.width = self._mm(w)
		self.borders(table, {})
		return table

	def _border_el(self, parent, edge, on, color=None, size=4):
		from docx.oxml import OxmlElement
		from docx.oxml.ns import qn

		el = OxmlElement(f"w:{edge}")
		if on:
			el.set(qn("w:val"), "single")
			el.set(qn("w:sz"), str(size))
			el.set(qn("w:space"), "0")
			el.set(qn("w:color"), _hex_attr(color or self.theme["border_color"]))
		else:
			el.set(qn("w:val"), "nil")
		parent.append(el)

	def borders(self, table, edges, color=None, size=4):
		"""edges: {top,bottom,left,right,insideH,insideV: bool}; missing edges are removed."""
		b = _ensure(table._tbl.tblPr, "w:tblBorders", replace=True)
		for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
			self._border_el(b, edge, edges.get(edge), color, size)

	def cell_borders(self, cell, edges, color=None, size=4):
		"""Borders on one cell; only the listed edges are set, the rest follow the table."""
		b = _ensure(cell._tc.get_or_add_tcPr(), "w:tcBorders", replace=True)
		for edge in ("top", "left", "bottom", "right"):
			if edge in edges:
				self._border_el(b, edge, edges[edge], color, size)

	def shade(self, cell, hex_color):
		from docx.oxml.ns import qn

		shd = _ensure(cell._tc.get_or_add_tcPr(), "w:shd", replace=True)
		shd.set(qn("w:val"), "clear")
		shd.set(qn("w:color"), "auto")
		shd.set(qn("w:fill"), _hex_attr(hex_color))

	def repeat_header(self, row):
		from docx.oxml.ns import qn

		_ensure(row._tr.get_or_add_trPr(), "w:tblHeader").set(qn("w:val"), "true")

	def rule(self, color=None, size=8):
		from docx.oxml import OxmlElement
		from docx.oxml.ns import qn
		from docx.shared import Pt

		p = self.doc.add_paragraph()
		p.paragraph_format.space_after = Pt(6)
		pBdr = _ensure(p._p.get_or_add_pPr(), "w:pBdr", replace=True)
		bottom = OxmlElement("w:bottom")
		bottom.set(qn("w:val"), "single")
		bottom.set(qn("w:sz"), str(size))
		bottom.set(qn("w:space"), "1")
		bottom.set(qn("w:color"), _hex_attr(color or self.theme["primary_color"]))
		pBdr.append(bottom)
		return p

	def spacer(self, pt=6):
		from docx.shared import Pt

		p = self.doc.add_paragraph()
		p.paragraph_format.space_after = Pt(0)
		p.paragraph_format.line_spacing = Pt(pt)
		return p

	def write_sections(self):
		for s in self.model["sections"]:
			if s.get("hidden"):
				continue
			if (s.get("style") or {}).get("page_break_before"):
				self.page_break()
			handler = getattr(self, f"s_{s['type']}", None)
			if handler:
				handler(s)

	def section_label(self, s):
		from docx.shared import Pt

		if not s.get("label"):
			return
		p = self.doc.add_paragraph()
		p.paragraph_format.space_before = Pt(8)
		p.paragraph_format.space_after = Pt(3)
		p.paragraph_format.keep_with_next = True
		self.run(
			p,
			s["label"],
			{"bold": True, "color": self.theme["primary_color"], "size": self.theme["base_font_size"] + 1},
		)

	def s_letterhead(self, s):
		from docx.enum.text import WD_ALIGN_PARAGRAPH
		from docx.shared import Mm

		company = s["company"]
		logo = _load_image(company.get("logo"))
		widths = [self.content_mm * 0.7, self.content_mm * 0.3] if logo else [self.content_mm]
		table = self.table(self.doc, 1, len(widths), widths)
		left = table.rows[0].cells[0]
		self.write_node(
			left,
			company["name"],
			{"size": self.theme["heading_font_size"] - 2, "color": self.theme["primary_color"]},
		)
		for line in company["lines"]:
			self.write_node(
				left,
				line,
				{"size": max(self.theme["base_font_size"] - 1, 7), "color": LABEL_COLOR},
				reuse=False,
			)
		if logo:
			p = table.rows[0].cells[1].paragraphs[0]
			p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
			p.add_run().add_picture(logo, height=Mm(18))
		self.rule()

	def s_title(self, s):
		from docx.shared import Pt

		p = self.doc.add_paragraph()
		p.paragraph_format.space_before = Pt(4)
		p.paragraph_format.space_after = Pt(2)
		base = {"size": self.theme["heading_font_size"], "color": self.theme["primary_color"]}
		align = (s["heading"].get("style") or {}).get("align") or self.theme["heading_align"].lower()
		self.align(p, align)
		if not s["heading"].get("hidden"):
			self.run(p, s["heading"]["display"], self.node_style(s["heading"], base))
		if s.get("subtitle") and not s["subtitle"].get("hidden"):
			self.write_node(
				self.doc, s["subtitle"], {"size": self.theme["base_font_size"] + 1}, reuse=False, align=align
			)
		self.spacer()

	def kv_rows(self, container, nodes, width_mm, label_ratio=0.38):
		"""Label / value pairs as a borderless 2-column table inside ``container``."""
		nodes = [n for n in nodes if not n.get("hidden")]
		if not nodes:
			return None
		widths = [width_mm * label_ratio, width_mm * (1 - label_ratio)]
		table = self.table(container, len(nodes), 2, widths)
		for row, node in zip(table.rows, nodes, strict=False):
			label_cell, value_cell = row.cells
			if node.get("label"):
				self.run(label_cell.paragraphs[0], node["label"], {"color": LABEL_COLOR})
			self.write_node(value_cell, node)
		return table

	def s_header(self, s):
		from docx.shared import Pt

		half = self.content_mm / 2
		table = self.table(self.doc, 1, 2, [half, half])
		left_cell, right_cell = table.rows[0].cells
		left = s.get("left")
		left_items = [n for n in (left or {}).get("items", []) if not n.get("hidden")]
		if left_items:
			self.borders(table, {"top": 1, "bottom": 1, "left": 1, "right": 1, "insideV": 1})
			p = left_cell.paragraphs[0]
			self.run(
				p,
				left["title"].upper(),
				{
					"bold": True,
					"size": self.theme["base_font_size"] - 1,
					"color": self.theme["primary_color"],
				},
			)
			p.paragraph_format.space_after = Pt(3)
			for node in left_items:
				label = node.get("label")
				if label:
					p = left_cell.add_paragraph()
					self.run(p, f"{label}: ", {"color": LABEL_COLOR})
					self.run(p, node["display"], self.node_style(node))
				else:
					self.write_node(left_cell, node, reuse=False)
		else:
			self.cell_borders(right_cell, {"top": 1, "bottom": 1, "left": 1, "right": 1})
		self.kv_rows(right_cell, s["right"]["items"], half - 4, label_ratio=0.35)
		# python-docx leaves an empty paragraph before a nested table; keep it tiny.
		right_cell.paragraphs[0].paragraph_format.space_after = Pt(0)
		self.spacer()

	def s_fields(self, s):
		columns = [[n for n in col if not n.get("hidden")] for col in s["columns"]]
		columns = [c for c in columns if c]
		if not columns:
			return
		self.section_label(s)
		if len(columns) == 1:
			self.kv_rows(self.doc, columns[0], self.content_mm, label_ratio=0.3)
		else:
			from docx.shared import Pt

			width = self.content_mm / len(columns)
			outer = self.table(self.doc, 1, len(columns), [width] * len(columns))
			if s.get("boxed", True):
				self.borders(outer, {"top": 1, "bottom": 1, "left": 1, "right": 1, "insideV": 1})
			for cell, nodes in zip(outer.rows[0].cells, columns, strict=False):
				self.kv_rows(cell, nodes, width - 4)
				cell.paragraphs[0].paragraph_format.space_after = Pt(0)
		self.spacer()

	def s_table(self, s):
		from docx.shared import Pt

		rows = s["rows"]
		cols = [c for c in s["columns"]]
		visible = [i for i, c in enumerate(cols) if not c.get("hidden")]
		if not rows or not visible:
			return
		self.section_label(s)
		sno_w = 10
		numeric = [i for i in visible if cols[i].get("align") == "right"]
		num_w = min(28, (self.content_mm - sno_w) * 0.18)
		text_cols = [i for i in visible if i not in numeric]
		text_w = (self.content_mm - sno_w - num_w * len(numeric)) / max(len(text_cols), 1)
		if not text_cols:
			num_w = (self.content_mm - sno_w) / len(numeric)
		widths = [sno_w] + [num_w if i in numeric else text_w for i in visible]

		table = self.table(self.doc, len(rows) + 1, len(widths), widths)
		# The table's own style (chosen in the preview) wins over the theme's.
		style = (s.get("table_style") or {}).get("table_style") or self.theme["table_style"]
		if style == "Bordered":
			self.borders(table, {e: 1 for e in ("top", "bottom", "left", "right", "insideH", "insideV")})
		else:
			self.borders(table, {"bottom": 1})

		head = table.rows[0]
		self.repeat_header(head)
		labels = ["S.No"] + [cols[i]["label"] for i in visible]
		aligns = ["center"] + [cols[i].get("align") or "left" for i in visible]
		for cell, label, align in zip(head.cells, labels, aligns, strict=False):
			if style != "Minimal":
				self.shade(cell, self.theme["header_shading"])
			else:
				self.cell_borders(cell, {"bottom": 1}, self.theme["primary_color"], 8)
			p = cell.paragraphs[0]
			self.align(p, align)
			self.run(p, label, {"bold": True})

		for n, (row, data) in enumerate(zip(table.rows[1:], rows, strict=False), 1):
			cells = row.cells
			p = cells[0].paragraphs[0]
			self.align(p, "center")
			self.run(p, str(n))
			for cell, i in zip(cells[1:], visible, strict=False):
				node = data["cells"][i]
				self.write_node(cell, node, align=cols[i].get("align"))
		for row in table.rows:
			for cell in row.cells:
				for p in cell.paragraphs:
					p.paragraph_format.space_after = Pt(1)

	def s_totals(self, s):
		from docx.enum.table import WD_TABLE_ALIGNMENT

		items = [n for n in s["items"] if not n.get("hidden")]
		if not items:
			return
		width = self.content_mm * 0.48
		table = self.table(self.doc, len(items), 2, [width * 0.55, width * 0.45])
		table.alignment = WD_TABLE_ALIGNMENT.RIGHT
		if self.theme["table_style"] == "Bordered":
			self.borders(table, {e: 1 for e in ("top", "bottom", "left", "right", "insideH", "insideV")})
		for row, node in zip(table.rows, items, strict=False):
			label_cell, value_cell = row.cells
			grand = node.get("grand")
			style = {"bold": True} if grand else {}
			p = label_cell.paragraphs[0]
			self.align(p, "right")
			self.run(p, node["label"], style)
			self.write_node(value_cell, node, style, align="right")
			if grand:
				for cell in row.cells:
					self.shade(cell, self.theme["header_shading"])
					self.cell_borders(cell, {"top": 1}, self.theme["primary_color"], 8)
		self.spacer()

	def s_rich_text(self, s):
		if s["node"].get("hidden"):
			return
		self.section_label(s)
		self.write_rich(self.doc, s["node"]["display"], s["node"].get("style"), reuse=False)
		self.spacer()

	s_terms = s_rich_text
	s_text = s_rich_text

	def s_signature(self, s):
		from docx.shared import Pt

		lines = [n for n in s["lines"] if not n.get("hidden")]
		if not lines:
			return
		self.spacer(12)
		for i, node in enumerate(lines):
			p = self.doc.add_paragraph()
			p.paragraph_format.keep_with_next = i < len(lines) - 1
			self.align(p, (node.get("style") or {}).get("align") or "right")
			if i == len(lines) - 1 and len(lines) > 1:
				p.paragraph_format.space_before = Pt(36)
			self.run(p, node["display"], self.node_style(node))


SEQ = {
	"w:rPr": (
		"w:rStyle",
		"w:rFonts",
		"w:b",
		"w:bCs",
		"w:i",
		"w:iCs",
		"w:caps",
		"w:smallCaps",
		"w:strike",
		"w:dstrike",
		"w:outline",
		"w:shadow",
		"w:emboss",
		"w:imprint",
		"w:noProof",
		"w:snapToGrid",
		"w:vanish",
		"w:webHidden",
		"w:color",
		"w:spacing",
		"w:w",
		"w:kern",
		"w:position",
		"w:sz",
		"w:szCs",
		"w:highlight",
		"w:u",
		"w:effect",
		"w:bdr",
		"w:shd",
		"w:fitText",
		"w:vertAlign",
		"w:rtl",
		"w:cs",
		"w:em",
		"w:lang",
		"w:eastAsianLayout",
		"w:specVanish",
		"w:oMath",
	),
	"w:tblPr": (
		"w:tblStyle",
		"w:tblpPr",
		"w:tblOverlap",
		"w:bidiVisual",
		"w:tblStyleRowBandSize",
		"w:tblStyleColBandSize",
		"w:tblW",
		"w:jc",
		"w:tblCellSpacing",
		"w:tblInd",
		"w:tblBorders",
		"w:shd",
		"w:tblLayout",
		"w:tblCellMar",
		"w:tblLook",
		"w:tblCaption",
		"w:tblDescription",
		"w:tblPrChange",
	),
	"w:tcPr": (
		"w:cnfStyle",
		"w:tcW",
		"w:gridSpan",
		"w:hMerge",
		"w:vMerge",
		"w:tcBorders",
		"w:shd",
		"w:noWrap",
		"w:tcMar",
		"w:textDirection",
		"w:tcFitText",
		"w:vAlign",
		"w:hideMark",
		"w:headers",
		"w:cellIns",
		"w:cellDel",
		"w:cellMerge",
		"w:tcPrChange",
	),
	"w:trPr": (
		"w:cnfStyle",
		"w:divId",
		"w:gridBefore",
		"w:gridAfter",
		"w:wBefore",
		"w:wAfter",
		"w:cantSplit",
		"w:trHeight",
		"w:tblHeader",
		"w:tblCellSpacing",
		"w:jc",
		"w:hidden",
		"w:ins",
		"w:del",
		"w:trPrChange",
	),
	"w:pPr": (
		"w:pStyle",
		"w:keepNext",
		"w:keepLines",
		"w:pageBreakBefore",
		"w:framePr",
		"w:widowControl",
		"w:numPr",
		"w:suppressLineNumbers",
		"w:pBdr",
		"w:shd",
		"w:tabs",
		"w:suppressAutoHyphens",
		"w:kinsoku",
		"w:wordWrap",
		"w:overflowPunct",
		"w:topLinePunct",
		"w:autoSpaceDE",
		"w:autoSpaceDN",
		"w:bidi",
		"w:adjustRightInd",
		"w:snapToGrid",
		"w:spacing",
		"w:ind",
		"w:contextualSpacing",
		"w:mirrorIndents",
		"w:suppressOverlap",
		"w:jc",
		"w:textDirection",
		"w:textAlignment",
		"w:textboxTightWrap",
		"w:outlineLvl",
		"w:divId",
		"w:cnfStyle",
		"w:rPr",
		"w:sectPr",
		"w:pPrChange",
	),
}


def _ensure(parent, tag, replace=False):
	"""Get (or replace) a child of a *Pr element, inserted at its schema position."""
	from docx.oxml import OxmlElement
	from docx.oxml.ns import qn

	el = parent.find(qn(tag))
	if el is not None and replace:
		parent.remove(el)
		el = None
	if el is None:
		el = OxmlElement(tag)
		parent_tag = "w:" + parent.tag.split("}")[1]
		seq = SEQ[parent_tag]
		parent.insert_element_before(el, *seq[seq.index(tag) + 1 :])
	return el


def _load_image(url):
	"""A PNG/JPEG stream for a site file or asset URL, or None."""
	if not url or url.startswith(("http://", "https://", "data:")):
		return None
	path = None
	try:
		if url.startswith(("/files/", "/private/files/")):
			name = frappe.db.get_value("File", {"file_url": url}, "name")
			if name:
				path = frappe.get_doc("File", name).get_full_path()
			else:
				path = frappe.get_site_path(url.lstrip("/") if url.startswith("/private") else "public" + url)
		elif url.startswith("/assets/"):
			path = os.path.join(frappe.local.sites_path, url.lstrip("/"))
		if not path or not os.path.exists(path):
			return None
		from PIL import Image

		img = Image.open(path)
		if img.mode not in ("RGB", "L"):
			rgba = img.convert("RGBA")
			bg = Image.new("RGB", rgba.size, (255, 255, 255))
			bg.paste(rgba, mask=rgba.split()[-1])
			img = bg
		out = BytesIO()
		img.save(out, format="PNG")
		out.seek(0)
		return out
	except Exception:
		frappe.log_error(title="DocStudio: could not load logo")
		return None
